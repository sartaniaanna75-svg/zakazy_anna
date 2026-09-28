"""Программа сопоставления и заказа. Открывается в браузере на этом компьютере."""

import sys
import threading
import webbrowser
from datetime import date
from pathlib import Path

from flask import Flask, flash, redirect, render_template, request, url_for
from openpyxl import Workbook, load_workbook
from werkzeug.utils import secure_filename

import convert_prices
import store

ROOT = Path(__file__).resolve().parents[1]
INBOX = ROOT / "inbox"
PRICE_DIR = Path(r"C:\Users\1\Desktop\Прайсы")
ORDERS_DIR = ROOT / "out" / "orders"

app = Flask(__name__)
app.secret_key = "zakazy-local"


def latest_offers(rows: list[dict]) -> list[dict]:
    chosen = {}
    for row in rows:
        key = (store.offer_key(row), row.get("Тип цены") or "", row.get("Исходный файл") or "")
        previous = chosen.get(key)
        if previous is None or store.parse_row_date(row.get("Дата обновления")) >= store.parse_row_date(previous.get("Дата обновления")):
            chosen[key] = row
    return list(chosen.values())


def auto_link() -> int:
    catalog = {item["barcode"]: item for item in store.load_catalog() if item.get("barcode")}
    matches = store.load_matches()
    added = 0
    seen = set()
    for row in store.load_prices():
        barcode = row.get("Штрихкод") or ""
        key = store.offer_key(row)
        if not barcode or barcode not in catalog or key in seen:
            continue
        seen.add(key)
        current = matches.get(key)
        if current and current.get("status") in {"confirmed", "rejected"}:
            continue
        item = catalog[barcode]
        problem = store.unit_problem(item.get("unit") or "", row.get("Единица измерения") or "")
        matches[key] = {
            "supplier": row.get("Поставщик") or "",
            "supplier_name": row.get("Наименование") or "",
            "supplier_code": row.get("Код товара") or "",
            "barcode": barcode,
            "catalog_code": item["code"],
            "method": "barcode",
            "status": "review",
            "problem": problem,
        }
        added += 1
    store.save_matches(matches)
    return added


def linked_rows():
    catalog = {item["code"]: item for item in store.load_catalog()}
    matches = store.load_matches()
    grouped = {}
    for row in latest_offers(store.load_prices()):
        match = matches.get(store.offer_key(row))
        if not match or match.get("status") == "rejected":
            continue
        item = catalog.get(match.get("catalog_code"))
        if not item:
            continue
        problem = store.unit_problem(item.get("unit") or "", row.get("Единица измерения") or "")
        bucket = grouped.setdefault(item["code"], {"item": item, "offers": [], "blocked": []})
        if problem:
            bucket["blocked"].append({**row, "problem": problem})
        else:
            bucket["offers"].append(row)
    return grouped


def build_order_lines():
    decisions = store.load_decisions()
    today = date.today().isoformat()
    absent = decisions["absent"]
    confirmed = decisions["confirmed"]
    lines = []
    blocked = []
    for code, bucket in linked_rows().items():
        item = bucket["item"]
        blocked.extend(
            {
                "code": code,
                "name": item["name"],
                "supplier": row.get("Поставщик"),
                "supplier_name": row.get("Наименование"),
                "problem": row["problem"],
            }
            for row in bucket["blocked"]
        )
        usable = []
        for row in bucket["offers"]:
            mark = f"{code}|{row.get('Поставщик')}"
            if absent.get(mark) == today:
                continue
            usable.append(row)
        if not usable and not bucket["blocked"]:
            continue
        units = {store.norm_unit(row.get("Единица измерения") or "") for row in usable}
        if len(units) > 1:
            blocked.append({"code": code, "name": item["name"], "supplier": "", "supplier_name": "", "problem": "единица не совпала"})
            continue
        fixed = confirmed.get(code)
        winner = None
        if fixed and fixed.get("date") == today:
            winner = next((row for row in usable if row.get("Поставщик") == fixed.get("supplier")), None)
            if winner is None:
                saved = fixed.get("row") or {}
                winner = {
                    "Поставщик": saved.get("supplier") or fixed.get("supplier"),
                    "Наименование": saved.get("supplier_name") or "",
                    "Цена": saved.get("price") or fixed.get("price"),
                    "Тип цены": saved.get("price_type") or "",
                    "Дата обновления": saved.get("updated") or "",
                    "Единица измерения": saved.get("unit") or "",
                }
        if winner is None and usable:
            winner = min(usable, key=lambda row: row.get("Цена") or 10**12)
        if winner is None:
            continue
        others = sorted((row for row in usable if row is not winner), key=lambda row: row.get("Цена") or 10**12)
        lines.append(
            {
                "code": code,
                "name": item["name"],
                "unit": item.get("unit") or (winner.get("Единица измерения") if isinstance(winner, dict) else ""),
                "supplier": winner.get("Поставщик") if isinstance(winner, dict) else fixed.get("supplier"),
                "supplier_name": winner.get("Наименование") if isinstance(winner, dict) else "",
                "price": winner.get("Цена") if isinstance(winner, dict) else fixed.get("price"),
                "price_type": winner.get("Тип цены") if isinstance(winner, dict) else "",
                "updated": winner.get("Дата обновления") if isinstance(winner, dict) else "",
                "confirmed": bool(fixed and fixed.get("date") == today and fixed.get("supplier") == (winner.get("Поставщик") if isinstance(winner, dict) else "")),
                "next_supplier": others[0].get("Поставщик") if others else "",
                "next_price": others[0].get("Цена") if others else "",
            }
        )
    lines.sort(key=lambda line: line["name"])
    return lines, blocked


def read_catalog_file(path: Path) -> list[dict]:
    if path.suffix.lower() == ".xls":
        import xlrd

        book = xlrd.open_workbook(path)
        sheet = book.sheet_by_index(0)
        rows = [sheet.row_values(index) for index in range(sheet.nrows)]
    else:
        book = load_workbook(path, data_only=True, read_only=True)
        sheet = book.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        book.close()
    header_index = None
    for index, row in enumerate(rows[:15]):
        labels = " ".join(convert_prices.as_text(cell).lower() for cell in row)
        if "наименование" in labels or "номенклатура" in labels:
            header_index = index
            break
    if header_index is None:
        raise ValueError("В файле нет колонки «наименование».")
    labels = [convert_prices.as_text(cell).lower() for cell in rows[header_index]]

    def find(*needles):
        for index, label in enumerate(labels):
            if any(needle in label for needle in needles):
                return index
        return None

    name_col = find("наименование", "номенклатура")
    code_col = find("код")
    barcode_col = find("штрихкод")
    unit_col = find("единица", "ед")
    if name_col is None:
        raise ValueError("В файле нет колонки «наименование».")
    items = []
    for row in rows[header_index + 1 :]:
        name = convert_prices.as_text(row[name_col] if name_col < len(row) else "")
        if not name:
            continue
        code = convert_prices.as_text(row[code_col] if code_col is not None and code_col < len(row) else "") or name
        items.append(
            {
                "code": code,
                "barcode": convert_prices.barcode_text(row[barcode_col]) if barcode_col is not None and barcode_col < len(row) else "",
                "name": name,
                "unit": convert_prices.as_text(row[unit_col] if unit_col is not None and unit_col < len(row) else ""),
            }
        )
    return items


@app.route("/")
def today_page():
    prices = store.load_prices()
    matches = store.load_matches()
    keys = {store.offer_key(row) for row in prices}
    linked = sum(1 for key in keys if matches.get(key, {}).get("status") in {"review", "confirmed"})
    unmatched = len(keys) - linked
    return render_template(
        "today.html",
        price_rows=len(prices),
        files=len({row.get("Исходный файл") for row in prices}),
        linked=linked,
        unmatched=unmatched,
        catalog=len(store.load_catalog()),
        can_order=linked > 0,
    )


@app.route("/prices", methods=["GET", "POST"])
def prices_page():
    if request.method == "POST":
        rows, notes, fresh = convert_prices.convert_folders([INBOX, PRICE_DIR], store.load_prices())
        store.save_prices(rows)
        convert_prices.write_table(rows, ROOT / "out" / "prices.xlsx")
        added = auto_link()
        (ROOT / "out" / "prices.log").write_text("\n".join(notes) + f"\n\nНовых строк: {len(fresh)}\n", encoding="utf-8")
        flash(f"Разобрано. Новых строк: {len(fresh)}. Связано по штрихкоду: {added}.")
        return render_template("prices.html", notes=notes)
    log_path = ROOT / "out" / "prices.log"
    notes = log_path.read_text(encoding="utf-8").splitlines() if log_path.exists() else []
    return render_template("prices.html", notes=notes)


@app.route("/nomenclature", methods=["GET", "POST"])
def nomenclature_page():
    if request.method == "POST":
        action = request.form.get("action")
        matches = store.load_matches()
        if action == "upload":
            upload = request.files.get("catalog")
            if not upload or not upload.filename:
                flash("Выберите файл справочника.")
            else:
                INBOX.mkdir(parents=True, exist_ok=True)
                target = INBOX / secure_filename(upload.filename)
                upload.save(target)
                try:
                    store.save_catalog(read_catalog_file(target))
                    count = auto_link()
                    flash(f"Справочник загружен: {len(store.load_catalog())} товаров. По штрихкоду предложено связей: {count}.")
                except Exception as error:
                    flash(str(error))
        elif action == "link":
            catalog_code = request.form.get("catalog_code") or ""
            offer = request.form.get("offer_key") or ""
            catalog = {item["code"]: item for item in store.load_catalog()}
            row = next((item for item in store.load_prices() if store.offer_key(item) == offer), None)
            if catalog_code not in catalog or row is None:
                flash("Выберите ваш товар слева и название поставщика справа.")
            else:
                item = catalog[catalog_code]
                matches[offer] = {
                    "supplier": row.get("Поставщик") or "",
                    "supplier_name": row.get("Наименование") or "",
                    "supplier_code": row.get("Код товара") or "",
                    "barcode": row.get("Штрихкод") or "",
                    "catalog_code": catalog_code,
                    "method": "manual",
                    "status": "confirmed",
                    "problem": store.unit_problem(item.get("unit") or "", row.get("Единица измерения") or ""),
                }
                store.save_matches(matches)
                flash("Связано. В следующий прайс этого поставщика программа узнает товар сама.")
        elif action in {"accept", "reject"}:
            offer = request.form.get("offer_key") or ""
            if offer in matches:
                matches[offer]["status"] = "confirmed" if action == "accept" else "rejected"
                store.save_matches(matches)
                flash("Связка подтверждена." if action == "accept" else "Это не тот товар. Строка в заказ не попадёт.")
        return redirect(url_for("nomenclature_page", q=request.form.get("q") or "", supplier_q=request.form.get("supplier_q") or ""))

    query = (request.args.get("q") or "").strip().lower()
    supplier_query = (request.args.get("supplier_q") or "").strip().lower()
    catalog = store.load_catalog()
    if query:
        catalog = [item for item in catalog if query in item["name"].lower() or query in item["code"].lower() or query in item.get("barcode", "")]
    matches = store.load_matches()
    review = []
    unmatched = []
    seen = set()
    for row in store.load_prices():
        key = store.offer_key(row)
        if key in seen:
            continue
        seen.add(key)
        match = matches.get(key)
        if match and match.get("status") == "review":
            review.append({**row, "key": key, "problem": match.get("problem") or ""})
        elif not match or match.get("status") == "rejected":
            if match and match.get("status") == "rejected":
                continue
            name = (row.get("Наименование") or "").lower()
            supplier = (row.get("Поставщик") or "").lower()
            if supplier_query and supplier_query not in name and supplier_query not in supplier:
                continue
            unmatched.append({**row, "key": key})
    return render_template(
        "nomenclature.html",
        catalog=catalog[:40],
        catalog_total=len(store.load_catalog()),
        review=review[:30],
        review_total=len(review),
        unmatched=unmatched[:40],
        unmatched_total=len(unmatched),
        q=request.args.get("q") or "",
        supplier_q=request.args.get("supplier_q") or "",
    )


@app.route("/order", methods=["GET", "POST"])
def order_page():
    decisions = store.load_decisions()
    today = date.today().isoformat()
    if request.method == "POST":
        action = request.form.get("action")
        code = request.form.get("code") or ""
        supplier = request.form.get("supplier") or ""
        if action == "absent":
            decisions["absent"][f"{code}|{supplier}"] = today
            if decisions["confirmed"].get(code, {}).get("supplier") == supplier:
                decisions["confirmed"].pop(code, None)
            store.save_decisions(decisions)
            flash(f"До конца дня {supplier} для этой позиции не участвует. Берём следующего по цене.")
        elif action == "confirm":
            lines, _blocked = build_order_lines()
            line = next((item for item in lines if item["code"] == code), None)
            if line:
                decisions["confirmed"][code] = {"date": today, "supplier": line["supplier"], "price": line["price"], "row": line}
                store.save_decisions(decisions)
                flash("Строка подтверждена. Новый прайс её не перекинет.")
        elif action == "export":
            lines, _blocked = build_order_lines()
            export_orders(lines)
            flash(f"Файлы заказов лежат в папке {ORDERS_DIR}")
        return redirect(url_for("order_page"))
    lines, blocked = build_order_lines()
    return render_template("order.html", lines=lines[:80], line_total=len(lines), blocked=blocked[:40], can_order=bool(store.load_matches()))


def export_orders(lines: list[dict]) -> None:
    ORDERS_DIR.mkdir(parents=True, exist_ok=True)
    grouped = {}
    for line in lines:
        grouped.setdefault(line["supplier"] or "без поставщика", []).append(line)
    for supplier, items in grouped.items():
        book = Workbook()
        sheet = book.active
        sheet.title = "Заказ"
        sheet.append(["Код 1С", "Наименование", "Единица", "Наименование поставщика", "Цена", "Тип цены", "Дата прайса", "Поставщик"])
        for line in items:
            sheet.append([line["code"], line["name"], line["unit"], line["supplier_name"], line["price"], line["price_type"], line["updated"], supplier])
        safe = "".join(ch if ch.isalnum() or ch in " _-" else "_" for ch in supplier)[:40]
        book.save(ORDERS_DIR / f"{date.today().isoformat()}_{safe}.xlsx")


if __name__ == "__main__":
    INBOX.mkdir(parents=True, exist_ok=True)
    if "--no-browser" not in sys.argv:
        threading.Timer(1.0, lambda: webbrowser.open("http://127.0.0.1:8765")).start()
    app.run(host="127.0.0.1", port=8765, debug=False)
